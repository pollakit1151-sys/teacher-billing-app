# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import os
import re
import sys
import datetime
import threading
import shutil

DESKTOP_DIR = r"C:\Users\Legion\Desktop\รอบบ่าย"
BACKUP_DIR = os.path.join(os.path.dirname(__file__), "master_templates_backup")
LOCAL_DIR = os.path.dirname(__file__)

THAI_MONTHS_FULL = [
    "", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
    "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"
]
THAI_MONTHS_SHORT = [
    "", "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
    "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."
]

FILE_DEFS = [
    {"idx": 0, "filename": "ช่างยนต์ - ครูประจำ ปวช..xlsx", "dept": "ช่างยนต์", "group_code": "reg_vc", "group_name": "ครูประจำ", "level": "ปวช."},
    {"idx": 1, "filename": "ช่างยนต์ - ครูประจำ ปวส..xlsx", "dept": "ช่างยนต์", "group_code": "reg_vs", "group_name": "ครูประจำ", "level": "ปวส."},
    {"idx": 2, "filename": "ช่างยนต์ - ครูพิเศษ ปวช..xlsx", "dept": "ช่างยนต์", "group_code": "sp_vc", "group_name": "ครูพิเศษ", "level": "ปวช."},
    {"idx": 3, "filename": "ช่างยนต์ - ครูพิเศษ ปวส..xlsx", "dept": "ช่างยนต์", "group_code": "sp_vs", "group_name": "ครูพิเศษ", "level": "ปวส."},
    {"idx": 4, "filename": "ยานยนต์ไฟฟ้า - ครูประจำ ปวช..xlsx", "dept": "ยานยนต์ไฟฟ้า", "group_code": "reg_vc", "group_name": "ครูประจำ", "level": "ปวช."},
    {"idx": 5, "filename": "ยานยนต์ไฟฟ้า - ครูประจำ ปวส..xlsx", "dept": "ยานยนต์ไฟฟ้า", "group_code": "reg_vs", "group_name": "ครูประจำ", "level": "ปวส."},
    {"idx": 6, "filename": "ยานยนต์ไฟฟ้า - ครูพิเศษ ปวช..xlsx", "dept": "ยานยนต์ไฟฟ้า", "group_code": "sp_vc", "group_name": "ครูพิเศษ", "level": "ปวช."},
    {"idx": 7, "filename": "ยานยนต์ไฟฟ้า - ครูพิเศษ ปวส..xlsx", "dept": "ยานยนต์ไฟฟ้า", "group_code": "sp_vs", "group_name": "ครูพิเศษ", "level": "ปวส."},
]

def safe_set_cell(ws, r, c, val):
    cell = ws.cell(r, c)
    if type(cell).__name__ == 'MergedCell':
        for rng in ws.merged_cells.ranges:
            if c in range(rng.min_col, rng.max_col + 1) and r in range(rng.min_row, rng.max_row + 1):
                ws.cell(rng.min_row, rng.min_col).value = val
                return
    else:
        cell.value = val

def sync_week_sheets(wb, round_weeks):
    """
    Ensure wb has exactly the weekly sheets corresponding to round_weeks.
    Example: round_weeks = [1, 2, 3, 4, 5] -> ['สัปดาห์ที่ 1', ..., 'สัปดาห์ที่ 5']
    Creates missing tabs, removes extraneous tabs, and preserves non-weekly tabs.
    """
    target_names = [f"สัปดาห์ที่ {w}" for w in round_weeks]
    existing_week_sheets = [s for s in wb.sheetnames if s.strip().startswith("สัปดาห์ที่")]
    
    # 1. Delete week sheets that are not in target_names
    for s in existing_week_sheets:
        if s not in target_names:
            del wb[s]
            
    # 2. Add missing week sheets by copying an existing week sheet
    current_weeks = [s for s in wb.sheetnames if s.strip().startswith("สัปดาห์ที่")]
    base_sheet = wb[current_weeks[0]] if current_weeks else wb.worksheets[0]
    
    for target in target_names:
        if target not in wb.sheetnames:
            new_ws = wb.copy_worksheet(base_sheet)
            new_ws.title = target

    # 3. Reorder sheets: [target_names..., other_sheets...]
    other_sheets = [s for s in wb.sheetnames if s not in target_names]
    desired_order = target_names + other_sheets
    
    sheet_dict = {ws.title: ws for ws in wb._sheets}
    wb._sheets = [sheet_dict[name] for name in desired_order if name in sheet_dict]

def get_week_dates(week_num, base_monday_str="2026-09-14", year="2569"):
    """
    Returns date dict with start/end day, month name, year, and day strings.
    """
    try:
        parts = [int(p) for p in base_monday_str.split('-')]
        base_m = datetime.date(parts[0], parts[1], parts[2])
    except Exception:
        base_m = datetime.date(2026, 9, 14)
        
    mon = base_m + datetime.timedelta(days=(week_num - 1) * 7)
    fri = mon + datetime.timedelta(days=4)
    
    y_int = int(year) if str(year).isdigit() else 2569
    y_short = str(y_int % 100)
    
    day_date_strs = []
    for i in range(5):
        d = mon + datetime.timedelta(days=i)
        m_short = THAI_MONTHS_SHORT[d.month] if d.month < len(THAI_MONTHS_SHORT) else ""
        day_date_strs.append(f"{d.day} {m_short}{y_short}")
        
    m_start_full = THAI_MONTHS_FULL[mon.month] if mon.month < len(THAI_MONTHS_FULL) else ""
    m_end_full = THAI_MONTHS_FULL[fri.month] if fri.month < len(THAI_MONTHS_FULL) else ""
    
    return {
        "start_day": mon.day,
        "start_month": m_start_full,
        "end_day": fri.day,
        "end_month": m_end_full,
        "year": str(y_int),
        "day_date_strs": day_date_strs,
        "range_str": f"{mon.day} {m_start_full} {y_int} - {fri.day} {m_end_full} {y_int}"
    }

def calculate_excel_signing_date(week_num, base_monday_str="2026-09-14", year="2569", week_holiday_map=None):
    """
    คำนวณวันเซ็นเอกสาร (ปกติวันจันทร์ถัดไป แต่ถ้าวันจันทร์เป็นวันหยุดให้เลื่อนเป็นวันอังคาร ถ้าอังคารหยุดเลื่อนเป็นวันพุธ ไล่ไปเรื่อยๆ)
    """
    try:
        parts = [int(p) for p in base_monday_str.split('-')]
        base_m = datetime.date(parts[0], parts[1], parts[2])
    except Exception:
        base_m = datetime.date(2026, 9, 14)
        
    w_int = int(week_num) if str(week_num).isdigit() else 1
    # Next Monday = base_m + timedelta(days=w_int * 7)
    cur_d = base_m + datetime.timedelta(days=w_int * 7)
    
    day_names = ['จันทร์', 'อังคาร', 'พุธ', 'พฤหัส', 'ศุกร์', 'เสาร์', 'อาทิตย์']
    
    for step in range(14):
        # Weekend: Saturday is 5, Sunday is 6 in python weekday()
        if cur_d.weekday() in (5, 6):
            cur_d += datetime.timedelta(days=1)
            continue
            
        w_at_d = 1 + (cur_d - base_m).days // 7
        d_holidays = []
        if week_holiday_map and isinstance(week_holiday_map, dict):
            raw_h = week_holiday_map.get(str(w_at_d)) or week_holiday_map.get(w_at_d) or []
            d_holidays = [str(h).strip() for h in raw_h]

        dow_name = day_names[cur_d.weekday()]
        is_holi = any(dow_name in h or h in dow_name for h in d_holidays)
        if is_holi:
            cur_d += datetime.timedelta(days=1)
            continue
            
        # Found first working day
        break
        
    y_base = int(year) if str(year).isdigit() else (cur_d.year + 543)
    y_thai = y_base if y_base > 2500 else (cur_d.year + 543)
    m_full = THAI_MONTHS_FULL[cur_d.month] if cur_d.month < len(THAI_MONTHS_FULL) else ""
    return f"วันที่ {cur_d.day} เดือน {m_full}  พ.ศ.{y_thai}"

def populate_weekly_sheet_teachers(ws, teachers, week_num, date_info=None, dept="ช่างยนต์", week_holiday_map=None, base_monday_str="2026-09-14", year="2569"):
    # Master week number in F7
    safe_set_cell(ws, 7, 6, week_num)
    
    # Top date range
    if date_info:
        safe_set_cell(ws, 3, 4, date_info.get('start_day', ''))
        safe_set_cell(ws, 3, 6, date_info.get('start_month', ''))
        safe_set_cell(ws, 3, 8, f"พ.ศ. {date_info.get('year', '2569')}")
        safe_set_cell(ws, 3, 10, date_info.get('end_day', ''))
        safe_set_cell(ws, 3, 12, date_info.get('end_month', ''))
        safe_set_cell(ws, 3, 14, f"พ.ศ. {date_info.get('year', '2569')}")
        
    day_names = ['จันทร์', 'อังคาร', 'พุธ', 'พฤหัส', 'ศุกร์']
    day_date_strs = date_info.get('day_date_strs', []) if date_info else []

    cur_w_holidays = []
    if week_holiday_map and isinstance(week_holiday_map, dict):
        raw_hw = week_holiday_map.get(str(week_num)) or week_holiday_map.get(week_num) or []
        cur_w_holidays = [str(h).strip() for h in raw_hw]

    for b_idx, teacher in enumerate(teachers):
        start_r = 7 + b_idx * 46
        g_idx = b_idx + 1
        name = teacher.get('name', '')
        pos = teacher.get('position', 'ครู')
        duty = teacher.get('duty', '')
        level = teacher.get('level', 'ปวช.')
        req_min = teacher.get('required_min', 12 if level == 'ปวช.' else 10)
        
        # Row 1..4 (top header for this teacher)
        safe_set_cell(ws, start_r - 5, 1, f"แผนกวิชา{dept}")
        safe_set_cell(ws, start_r - 3, 7, "ü")
        safe_set_cell(ws, start_r - 3, 8, pos)
        safe_set_cell(ws, start_r - 3, 9, "หน้าที่พิเศษ")
        safe_set_cell(ws, start_r - 3, 10, duty if duty else "")
        
        # Row start_r (Main table header for teacher)
        safe_set_cell(ws, start_r, 1, g_idx)
        safe_set_cell(ws, start_r, 2, f"=AB{3 + b_idx}")  # Links to right summary
        safe_set_cell(ws, start_r, 5, pos)
        safe_set_cell(ws, start_r, 6, week_num)
        
        # Right summary table (Row 3 + b_idx, Col Z to AH)
        r_sum = 3 + b_idx
        safe_set_cell(ws, r_sum, 26, g_idx)
        safe_set_cell(ws, r_sum, 28, name)
        claim_hrs = teacher.get('claim_hours', 0)
        total_money = teacher.get('total_money', 0)
        safe_set_cell(ws, r_sum, 31, claim_hrs)
        safe_set_cell(ws, r_sum, 34, total_money)
        
        # Classes (5 days x 5 slots)
        day_classes = [[], [], [], [], []]
        for c in teacher.get('classes', []):
            d_str = str(c.get('day', '')).strip()
            d_idx = -1
            if 'จันทร์' in d_str: d_idx = 0
            elif 'อังคาร' in d_str: d_idx = 1
            elif 'พุธ' in d_str: d_idx = 2
            elif 'พฤหัส' in d_str: d_idx = 3
            elif 'ศุกร์' in d_str: d_idx = 4
            if d_idx >= 0:
                day_classes[d_idx].append(c)

        PEACH_FILL = PatternFill(fill_type='solid', start_color='FCE4D6', end_color='FCE4D6')
        PASTEL_RED_FILL = PatternFill(fill_type='solid', start_color='F2DBDB', end_color='F2DBDB')
        PASTEL_YELLOW_FILL = PatternFill(fill_type='solid', start_color='FEF9C3', end_color='FEF9C3')
        RED_BOLD_FONT = Font(name='TH SarabunPSK', size=14, bold=True, color='FFFF0000')
        BOLD_FONT = Font(name='TH SarabunPSK', size=14, bold=True, color='FF000000')
        CENTER_ALIGN = Alignment(horizontal='center', vertical='center')

        for day_idx in range(5):
            d_name = day_names[day_idx]
            day_c_list = day_classes[day_idx]
            is_holiday_day = any(d_name in h or h in d_name for h in cur_w_holidays) or any(c.get('is_holiday') for c in day_c_list)
            absent_classes = [c for c in day_c_list if c.get('is_absent') or c.get('is_substituted')]
            is_absent_day = not is_holiday_day and len(absent_classes) > 0

            if is_holiday_day:
                for slot in range(5):
                    ro = day_idx * 5 + slot
                    r = start_r + ro

                    if slot == 0:
                        safe_set_cell(ws, r, 7, day_names[day_idx])
                        ws.cell(r, 7).font = BOLD_FONT
                        safe_set_cell(ws, r, 8, "หยุด")
                        cell_h = ws.cell(r, 8)
                        cell_h.font = BOLD_FONT
                        cell_h.alignment = CENTER_ALIGN
                    elif slot == 1:
                        if day_idx < len(day_date_strs):
                            safe_set_cell(ws, r, 7, day_date_strs[day_idx])
                            ws.cell(r, 7).font = BOLD_FONT
                        else:
                            safe_set_cell(ws, r, 7, "")
                        safe_set_cell(ws, r, 8, "")
                    else:
                        safe_set_cell(ws, r, 7, "")
                        safe_set_cell(ws, r, 8, "")

                    safe_set_cell(ws, r, 9, "")
                    safe_set_cell(ws, r, 10, "")
                    safe_set_cell(ws, r, 12, "")
                    safe_set_cell(ws, r, 13, "")
                    safe_set_cell(ws, r, 14, "")
                    safe_set_cell(ws, r, 15, f"=M{r}*N{r}")
                    safe_set_cell(ws, r, 16, "")
                    safe_set_cell(ws, r, 17, "")
                    safe_set_cell(ws, r, 18, "")
                    safe_set_cell(ws, r, 19, f"=Q{r}*R{r}")
                    safe_set_cell(ws, r, 20, "")
                    safe_set_cell(ws, r, 21, "")

                    # Apply pastel red fill only to columns 8..10 (รหัสวิชา, ชั้น แผนก, เวลาสอน)
                    for col_idx in range(8, 11):
                        ws.cell(r, col_idx).fill = PASTEL_RED_FILL
            elif is_absent_day:
                first_abs = absent_classes[0]
                leave_reason = first_abs.get('leave_reason') or first_abs.get('reason') or 'ไปราชการ'
                sub_by = first_abs.get('sub_by') or 'ครูสอนแทน'
                hrs_text = first_abs.get('sub_alloc_text') or f"{sub_by} สอนแทน (นอก 4 ชม.)"
                if not hrs_text.startswith(sub_by):
                    hrs_text = f"{sub_by} สอนแทน ({hrs_text})"

                for slot in range(5):
                    ro = day_idx * 5 + slot
                    r = start_r + ro

                    if slot == 0:
                        safe_set_cell(ws, r, 7, day_names[day_idx])
                        ws.cell(r, 7).font = BOLD_FONT
                        safe_set_cell(ws, r, 8, leave_reason)
                        cell_h = ws.cell(r, 8)
                        cell_h.font = BOLD_FONT
                        cell_h.fill = PEACH_FILL
                        cell_h.alignment = CENTER_ALIGN
                    elif slot == 1:
                        if day_idx < len(day_date_strs):
                            safe_set_cell(ws, r, 7, day_date_strs[day_idx])
                            ws.cell(r, 7).font = BOLD_FONT
                        else:
                            safe_set_cell(ws, r, 7, "")
                        safe_set_cell(ws, r, 8, hrs_text)
                        cell_h = ws.cell(r, 8)
                        cell_h.font = RED_BOLD_FONT
                        cell_h.alignment = CENTER_ALIGN
                    else:
                        safe_set_cell(ws, r, 7, "")
                        safe_set_cell(ws, r, 8, "")

                    safe_set_cell(ws, r, 9, "")
                    safe_set_cell(ws, r, 10, "")
                    safe_set_cell(ws, r, 12, "")
                    safe_set_cell(ws, r, 13, "")
                    safe_set_cell(ws, r, 14, "")
                    safe_set_cell(ws, r, 15, f"=M{r}*N{r}")
                    safe_set_cell(ws, r, 16, "")
                    safe_set_cell(ws, r, 17, "")
                    safe_set_cell(ws, r, 18, "")
                    safe_set_cell(ws, r, 19, f"=Q{r}*R{r}")
                    safe_set_cell(ws, r, 20, "")
                    safe_set_cell(ws, r, 21, "")
            else:
                for slot in range(5):
                    ro = day_idx * 5 + slot
                    r = start_r + ro
                    c = day_classes[day_idx][slot] if slot < len(day_classes[day_idx]) else {}

                    if slot == 0:
                        safe_set_cell(ws, r, 7, day_names[day_idx])
                        ws.cell(r, 7).font = BOLD_FONT
                    elif slot == 1 and day_idx < len(day_date_strs):
                        safe_set_cell(ws, r, 7, day_date_strs[day_idx])
                        ws.cell(r, 7).font = BOLD_FONT
                    else:
                        safe_set_cell(ws, r, 7, "")

                    code = c.get("code", "")
                    safe_set_cell(ws, r, 8, code if code else "")

                    c_info = str(c.get("class_info", "") or "").strip()
                    if c_info and not re.search(r'\(\s*\d+\s*\)$', c_info):
                        c_info = f"{c_info} (26)"
                    safe_set_cell(ws, r, 9, c_info if c_info else "")
                    safe_set_cell(ws, r, 10, c.get("time_str", "") or "")

                    in_vc = c.get("in_vc", 0)
                    out_vc = c.get("out_vc", 0)
                    in_vs = c.get("in_vs", 0)
                    out_vs = c.get("out_vs", 0)

                    safe_set_cell(ws, r, 12, in_vc if in_vc > 0 else "")
                    safe_set_cell(ws, r, 13, out_vc if out_vc > 0 else "")
                    safe_set_cell(ws, r, 14, 200 if out_vc > 0 else "")
                    safe_set_cell(ws, r, 15, f"=M{r}*N{r}" if out_vc > 0 else "")

                    safe_set_cell(ws, r, 16, in_vs if in_vs > 0 else "")
                    safe_set_cell(ws, r, 17, out_vs if out_vs > 0 else "")
                    safe_set_cell(ws, r, 18, 270 if out_vs > 0 else "")
                    safe_set_cell(ws, r, 19, f"=Q{r}*R{r}" if out_vs > 0 else "")

                    safe_set_cell(ws, r, 20, c.get("note", "") or "")
                    safe_set_cell(ws, r, 21, c.get("note2", "") or "")

                    # Format workplace class hour cells in pastel light yellow ONLY if cell has number
                    is_wp = c.get("is_workplace") or ('สถานประกอบการ' in str(c.get('room', '')) or 'สถานประกอบการ' in str(c.get('class_info', '')) or 'สถานประกอบการ' in str(c.get('subject_name', '')) or 'สถานประกอบการ' in str(c.get('name', '')))
                    if is_wp:
                        if in_vc > 0:
                            ws.cell(r, 12).fill = PASTEL_YELLOW_FILL
                        if out_vc > 0:
                            ws.cell(r, 13).fill = PASTEL_YELLOW_FILL
                        if in_vs > 0:
                            ws.cell(r, 16).fill = PASTEL_YELLOW_FILL
                        if out_vs > 0:
                            ws.cell(r, 17).fill = PASTEL_YELLOW_FILL

                    # Format substitute class
                    if c.get("is_substitute"):
                        abs_name = c.get("absent_teacher_name", "")
                        clean_abs = re.sub(r'^(นาย|นางสาว|นาง)\s*', '', abs_name).strip()
                        sub_label = c.get("sub_label") or f"สอนแทน อ.{clean_abs}"
                        # Shift substitute label down 1 row if it occurs at start_r (Monday slot 0), to avoid overwriting teacher name
                        sub_r = r + 1 if r == start_r else r
                        safe_set_cell(ws, sub_r, 2, sub_label)
                        cell_b = ws.cell(sub_r, 2)
                        cell_b.font = RED_BOLD_FONT
                        # Keep Column 2 clean white background without box fill

                        for col_idx in range(3, 20):
                            ws.cell(r, col_idx).fill = PEACH_FILL
                
        # Row 34: Quota
        r_quota = start_r + 27
        if level == 'ปวส.':
            safe_set_cell(ws, r_quota, 12, 0)
            safe_set_cell(ws, r_quota, 16, req_min)
        else:
            safe_set_cell(ws, r_quota, 12, req_min)
            safe_set_cell(ws, r_quota, 16, 0)
            
        # Signature
        safe_set_cell(ws, start_r + 32, 2, f"=B{start_r}")
        sign_date_str = calculate_excel_signing_date(week_num, base_monday_str=base_monday_str, year=year, week_holiday_map=week_holiday_map)
        safe_set_cell(ws, start_r + 34, 1, sign_date_str)

    # Master signature date cell A41
    master_sign_date_str = calculate_excel_signing_date(week_num, base_monday_str=base_monday_str, year=year, week_holiday_map=week_holiday_map)
    safe_set_cell(ws, 41, 1, master_sign_date_str)

    # Clear right summary table rows beyond len(teachers)
    for r in range(3 + len(teachers), 35):
        for c in [26, 27, 28, 29, 30, 31, 32, 33, 34]:
            safe_set_cell(ws, r, c, None)

    # Cut off unused blocks if len(teachers) > 0
    if len(teachers) > 0:
        cutoff_row = 1 + len(teachers) * 46
        merged_to_remove = [m for m in ws.merged_cells.ranges if m.min_row >= cutoff_row]
        for m in merged_to_remove:
            ws.merged_cells.remove(m)
        if ws.max_row >= cutoff_row:
            ws.delete_rows(cutoff_row, ws.max_row - cutoff_row + 1)
    else:
        cutoff_row = 47
        if ws.max_row >= cutoff_row:
            ws.delete_rows(cutoff_row, ws.max_row - cutoff_row + 1)

def update_individual_summary_sheet(ws, teachers, round_weeks, dept="ช่างยนต์", group_name="ครูประจำ"):
    safe_set_cell(ws, 2, 4, f"แผนก{dept}")
    if dept == "ยานยนต์ไฟฟ้า":
        safe_set_cell(ws, 1, 4, group_name)
    else:
        safe_set_cell(ws, 4, 4, group_name)
    
    is_ev = (dept == "ยานยนต์ไฟฟ้า")
    header_w_r = 4 if is_ev else 5
    header_sub_r = 5 if is_ev else 6
    start_r = 6 if is_ev else 7
    
    # Week numbers
    for w_idx, w_num in enumerate(round_weeks):
        col = 3 + w_idx * 2
        safe_set_cell(ws, header_w_r, col, w_num)
        safe_set_cell(ws, header_sub_r, col, "ชม/ส")
        safe_set_cell(ws, header_sub_r, col + 1, "เงิน")
        
    safe_set_cell(ws, header_w_r, 15, "รวม / ชม.")
    safe_set_cell(ws, header_w_r, 16, "รวม / เงิน")
    
    first_week_sheet = f"สัปดาห์ที่ {round_weeks[0]}" if round_weeks else "สัปดาห์ที่ 1"
    
    for idx, t in enumerate(teachers):
        r = start_r + idx
        g_idx = idx + 1
        safe_set_cell(ws, r, 1, g_idx)
        safe_set_cell(ws, r, 2, f"='{first_week_sheet}'!AB{3 + idx}")
        
        hour_cols = []
        money_cols = []
        for w_idx, w_num in enumerate(round_weeks):
            w_sheet = f"สัปดาห์ที่ {w_num}"
            c_hr = 3 + w_idx * 2
            c_mn = c_hr + 1
            hr_letter = openpyxl.utils.get_column_letter(c_hr)
            mn_letter = openpyxl.utils.get_column_letter(c_mn)
            hour_cols.append(f"{hr_letter}{r}")
            money_cols.append(f"{mn_letter}{r}")
            
            safe_set_cell(ws, r, c_hr, f"='{w_sheet}'!AE{3 + idx}")
            safe_set_cell(ws, r, c_mn, f"='{w_sheet}'!AH{3 + idx}")
            
        for extra_idx in range(len(round_weeks), 6):
            c_hr = 3 + extra_idx * 2
            safe_set_cell(ws, r, c_hr, 0)
            safe_set_cell(ws, r, c_hr + 1, 0)
            
        safe_set_cell(ws, r, 15, f"={'+'.join(hour_cols)}")
        safe_set_cell(ws, r, 16, f"={'+'.join(money_cols)}")

    # Clear remaining rows up to row 35
    for r in range(start_r + len(teachers), 36):
        for c in range(1, 17):
            safe_set_cell(ws, r, c, None)

def update_cover_sheet(ws, teachers, date_range="", dept="ช่างยนต์"):
    if date_range:
        safe_set_cell(ws, 2, 9, date_range)
        
    ind_start_r = 6 if dept == "ยานยนต์ไฟฟ้า" else 7
    
    # Page 1: teachers 0..7 (up to 8 teachers, rows 8..15)
    for idx in range(min(len(teachers), 8)):
        t = teachers[idx]
        r = 8 + idx
        g_idx = idx + 1
        level = t.get('level', 'ปวช.')
        pos = t.get('position', 'ครู')
        ind_r = ind_start_r + idx
        
        safe_set_cell(ws, r, 1, g_idx)
        safe_set_cell(ws, r, 2, f"=ตารางแจงเงินรายบุคคล!B{ind_r}")
        safe_set_cell(ws, r, 6, pos)
        if level == 'ปวส.':
            safe_set_cell(ws, r, 7, "")
            safe_set_cell(ws, r, 8, "ü")
        else:
            safe_set_cell(ws, r, 7, "ü")
            safe_set_cell(ws, r, 8, "")
            
        safe_set_cell(ws, r, 9, f"=ตารางแจงเงินรายบุคคล!O{ind_r}")
        safe_set_cell(ws, r, 10, f"=ตารางแจงเงินรายบุคคล!P{ind_r}")
        
    # Clear unused rows on Page 1
    for idx in range(len(teachers), 8):
        r = 8 + idx
        for c in [1, 2, 6, 7, 8, 9, 10]:
            safe_set_cell(ws, r, c, None)
            
    # Page 2: teachers 8..14 (up to 7 teachers, rows 35..41)
    if len(teachers) > 8:
        safe_set_cell(ws, 34, 2, "ย/ม")
        safe_set_cell(ws, 34, 10, "=J16")
        
        for idx in range(8, min(len(teachers), 15)):
            t = teachers[idx]
            r = 35 + (idx - 8)
            g_idx = idx + 1
            level = t.get('level', 'ปวช.')
            pos = t.get('position', 'ครู')
            ind_r = ind_start_r + idx
            
            safe_set_cell(ws, r, 1, g_idx)
            safe_set_cell(ws, r, 2, f"=ตารางแจงเงินรายบุคคล!B{ind_r}")
            safe_set_cell(ws, r, 6, pos)
            if level == 'ปวส.':
                safe_set_cell(ws, r, 7, "")
                safe_set_cell(ws, r, 8, "ü")
            else:
                safe_set_cell(ws, r, 7, "ü")
                safe_set_cell(ws, r, 8, "")
                
            safe_set_cell(ws, r, 9, f"=ตารางแจงเงินรายบุคคล!O{ind_r}")
            safe_set_cell(ws, r, 10, f"=ตารางแจงเงินรายบุคคล!P{ind_r}")
            
        # Clear unused rows on Page 2
        for idx in range(len(teachers), 15):
            r = 35 + (idx - 8)
            for c in [1, 2, 6, 7, 8, 9, 10]:
                safe_set_cell(ws, r, c, None)
    else:
        # Clear Page 2 teacher rows 35..41 if 8 or fewer teachers
        for r in range(35, 42):
            for c in [1, 2, 6, 7, 8, 9, 10]:
                safe_set_cell(ws, r, c, None)

def update_leave_summary_sheet(ws, leaves, substitutions):
    for r in range(2, min(ws.max_row or 1, 50)):
        for c in range(1, 15):
            safe_set_cell(ws, r, c, None)
            
    r_curr = 2
    for item in (leaves or []):
        safe_set_cell(ws, r_curr, 1, item.get('teacher_name', ''))
        safe_set_cell(ws, r_curr, 2, item.get('type', 'ลา'))
        safe_set_cell(ws, r_curr, 3, f"สัปดาห์ที่ {item.get('week', '')}")
        safe_set_cell(ws, r_curr, 4, item.get('date', ''))
        r_curr += 1
        
    for sub in (substitutions or []):
        safe_set_cell(ws, r_curr, 1, sub.get('original_teacher_name', ''))
        safe_set_cell(ws, r_curr, 2, "สอนแทน")
        safe_set_cell(ws, r_curr, 4, sub.get('substitute_teacher_name', ''))
        safe_set_cell(ws, r_curr, 6, sub.get('time_str', ''))
        r_curr += 1

def update_single_template_file(fdef, calculated_data, round_weeks, week_date_map=None, base_monday_str="2026-09-14", year="2569", date_range="", leaves=None, substitutions=None, week_holiday_map=None):
    fname = fdef["filename"]
    dept = fdef["dept"]
    gcode = fdef["group_code"]
    gname = fdef["group_name"]
    
    tpl_path = os.path.join(BACKUP_DIR, fname)
    if not os.path.exists(tpl_path):
        tpl_path = os.path.join(DESKTOP_DIR, fname)
    if not os.path.exists(tpl_path):
        print(f"Notice: template {fname} not found, skipping.")
        return False
        
    wb = openpyxl.load_workbook(tpl_path)
    
    # 1. Sync weekly tabs
    sync_week_sheets(wb, round_weeks)
    
    # 2. Filter teachers for this file
    file_teachers = [t for t in calculated_data if t.get('dept', 'ช่างยนต์') == dept and t.get('group_code') == gcode]
    
    # 3. Populate each weekly sheet
    for w in round_weeks:
        sheet_title = f"สัปดาห์ที่ {w}"
        if sheet_title in wb.sheetnames:
            ws = wb[sheet_title]
            d_info = get_week_dates(w, base_monday_str=base_monday_str, year=year)
            populate_weekly_sheet_teachers(
                ws,
                file_teachers,
                week_num=w,
                date_info=d_info,
                dept=dept,
                week_holiday_map=week_holiday_map,
                base_monday_str=base_monday_str,
                year=year
            )
            
    # 4. Update individual summary sheet
    if 'ตารางแจงเงินรายบุคคล' in wb.sheetnames:
        update_individual_summary_sheet(wb['ตารางแจงเงินรายบุคคล'], file_teachers, round_weeks, dept=dept, group_name=gname)
        
    # 5. Update cover sheet
    if 'งบหน้ารวมครู' in wb.sheetnames:
        update_cover_sheet(wb['งบหน้ารวมครู'], file_teachers, date_range=date_range, dept=dept)
        
    # 6. Update leave summary sheet
    for sname in wb.sheetnames:
        if 'สรุปลา' in sname:
            update_leave_summary_sheet(wb[sname], leaves or [], substitutions or [])
            
    # 7. Global scan: erase any remaining stale 'เปรม' or 'เพ็งยอด'
    for s_name in wb.sheetnames:
        ws_chk = wb[s_name]
        for row in ws_chk.iter_rows(max_row=min(ws_chk.max_row or 1, 150), max_col=min(ws_chk.max_column or 1, 35)):
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    if 'เปรม' in cell.value or 'เพ็งยอด' in cell.value:
                        cell.value = ""

    # 8. Save to Desktop directory
    os.makedirs(DESKTOP_DIR, exist_ok=True)
    out_path = os.path.join(DESKTOP_DIR, fname)
    wb.save(out_path)
    wb.close()
    return True

def update_all_8_desktop_files(calculated_data, round_weeks, week_date_map=None, base_monday_str="2026-09-14", year="2569", date_range="", leaves=None, substitutions=None, week_holiday_map=None):
    for fdef in FILE_DEFS:
        try:
            update_single_template_file(
                fdef=fdef,
                calculated_data=calculated_data,
                round_weeks=round_weeks,
                week_date_map=week_date_map,
                base_monday_str=base_monday_str,
                year=year,
                date_range=date_range,
                leaves=leaves,
                substitutions=substitutions,
                week_holiday_map=week_holiday_map
            )
        except Exception as e:
            print(f"Error updating {fdef['filename']}: {e}")

def export_to_excel(calculated_data, week_num=1, target_path=None, date_range="", term="2", year="2569", round_num=1, round_weeks=None, dept="ช่างยนต์", file_idx=None, leaves=None, substitutions=None, base_monday_str="2026-09-14", week_holiday_map=None):
    """
    Main export entry point:
    1. Updates the active file immediately for fast user download.
    2. Updates all 8 files in Desktop\\รอบบ่าย.
    3. Saves a download copy for the browser.
    """
    if round_weeks is None:
        round_weeks = [1, 2, 3, 4, 5]
        
    target_fdef = None
    if file_idx is not None and 0 <= file_idx < len(FILE_DEFS):
        target_fdef = FILE_DEFS[file_idx]
    else:
        target_fdef = next((f for f in FILE_DEFS if f['dept'] == dept), FILE_DEFS[0])
        
    # Update targeted file immediately
    update_single_template_file(
        fdef=target_fdef,
        calculated_data=calculated_data,
        round_weeks=round_weeks,
        base_monday_str=base_monday_str,
        year=year,
        date_range=date_range,
        leaves=leaves,
        substitutions=substitutions,
        week_holiday_map=week_holiday_map
    )
    
    active_saved_path = os.path.join(DESKTOP_DIR, target_fdef["filename"])
    
    download_fname = f"เบิกรอบบ่าย_{term}-{year}_รอบ{round_num}.xlsx"
    desktop_dl = os.path.join(r"C:\Users\Legion\Desktop", download_fname)
    local_dl = os.path.join(LOCAL_DIR, "เบิกรอบบ่าย.xlsx")
    
    try:
        shutil.copy2(active_saved_path, local_dl)
    except Exception as e:
        print(f"Notice: copy to local_dl failed: {e}")
        
    try:
        shutil.copy2(active_saved_path, desktop_dl)
    except Exception as e:
        print(f"Notice: copy to desktop_dl failed: {e}")
        
    # Background thread to update other files
    def _bg_update_rest():
        for fdef in FILE_DEFS:
            if fdef["idx"] != target_fdef["idx"]:
                try:
                    update_single_template_file(
                        fdef=fdef,
                        calculated_data=calculated_data,
                        round_weeks=round_weeks,
                        base_monday_str=base_monday_str,
                        year=year,
                        date_range=date_range,
                        leaves=leaves,
                        substitutions=substitutions,
                        week_holiday_map=week_holiday_map
                    )
                except Exception as ex:
                    print(f"Background update error on {fdef['filename']}: {ex}")
                    
    t = threading.Thread(target=_bg_update_rest, daemon=True)
    t.start()
    
    return True
